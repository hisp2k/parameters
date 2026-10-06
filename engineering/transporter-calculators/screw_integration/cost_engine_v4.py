"""Costing engine v4 for questionnaire-driven transporter calculator.

Purpose
-------
Turn the deterministic engineering BOM from v3 into an auditable preliminary
cost estimate.  The module intentionally separates:
1) engineering quantities (BOM),
2) price evidence (corporate / supplier / online / estimate),
3) internal labor & overhead assumptions,
4) commercial price.

No old production labor rate is hard-coded.  Internal rates must come from the
cost master or from an explicit manual override.
"""
from __future__ import annotations

from dataclasses import dataclass, asdict, field
from datetime import date, datetime
import json
import math
import re
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple


SOURCE_CONFIDENCE = {
    "corporate": 0.97,
    "supplier_quote": 0.93,
    "last_purchase": 0.90,
    "online_market": 0.75,
    "formula_estimate": 0.55,
    "manual": 0.85,
    "missing": 0.00,
}

SOURCE_PRIORITY = {name: i for i, name in enumerate(
    ["corporate", "supplier_quote", "last_purchase", "manual", "online_market", "formula_estimate"])}

SOURCE_UNCERTAINTY = {
    "corporate": 0.04,
    "supplier_quote": 0.06,
    "last_purchase": 0.08,
    "online_market": 0.15,
    "formula_estimate": 0.25,
    "manual": 0.12,
    "missing": 0.50,
}


@dataclass
class PriceEvidence:
    category: str
    description: str
    unit: str
    unit_price_rub: float
    source_type: str = "missing"
    source_name: str = ""
    source_url: str = ""
    price_date: str = ""
    note: str = ""
    confidence: float = 0.0
    item_key: str = ""

    def __post_init__(self) -> None:
        if self.confidence <= 0:
            self.confidence = SOURCE_CONFIDENCE.get(self.source_type, 0.0)


@dataclass
class LaborRate:
    operation: str
    rate_rub_h: float
    source: str = ""
    note: str = ""


@dataclass
class CostSettings:
    vat_pct: float = 22.0
    target_margin_pct: float = 20.0
    contingency_pct: float = 8.0
    procurement_overhead_pct: float = 3.0
    production_overhead_pct: float = 15.0
    packaging_pct: float = 1.5
    warranty_risk_pct: float = 2.0
    commissioning_reserve_pct: float = 2.0
    minimum_price_coverage_pct: float = 75.0
    frame_steel_waste_pct: float = 8.0
    belt_length_reserve_pct: float = 3.0
    installation_complexity_factor: float = 1.0


@dataclass
class CostMaster:
    settings: CostSettings = field(default_factory=CostSettings)
    labor_rates: Dict[str, LaborRate] = field(default_factory=dict)
    prices: List[PriceEvidence] = field(default_factory=list)
    norms: Dict[str, float] = field(default_factory=dict)

    def price_candidates(self, category: str) -> List[PriceEvidence]:
        return [p for p in self.prices if p.category.lower().strip() == category.lower().strip() and p.unit_price_rub > 0]


@dataclass
class PricedBOMItem:
    position: str
    quantity: float
    unit: str
    parameter: str
    category: str
    unit_price_rub: float
    total_rub: float
    source_type: str
    source_name: str
    source_url: str
    confidence: float
    uncertainty_pct: float
    status: str
    note: str = ""
    price_date: str = ""


@dataclass
class LaborLine:
    operation: str
    hours: float
    rate_rub_h: float
    cost_rub: float
    source: str
    status: str


@dataclass
class CostSummary:
    product_type: str
    direct_materials_rub: float
    direct_labor_rub: float
    procurement_overhead_rub: float
    production_overhead_rub: float
    contingency_rub: float
    packaging_rub: float
    warranty_risk_rub: float
    commissioning_reserve_rub: float
    full_cost_rub: float
    selling_price_ex_vat_rub: float
    vat_rub: float
    selling_price_inc_vat_rub: float
    price_low_inc_vat_rub: float
    price_high_inc_vat_rub: float
    price_coverage_pct: float
    weighted_confidence: float
    weighted_uncertainty_pct: float
    pricing_status: str
    warnings: List[str]
    priced_bom: List[Dict[str, Any]]
    labor: List[Dict[str, Any]]

    def to_json(self) -> str:
        return json.dumps(asdict(self), ensure_ascii=False, indent=2, default=str)


# ---------------------------------------------------------------------------
# Cost-master loading
# ---------------------------------------------------------------------------

def _f(v: Any, default: float = 0.0) -> float:
    if v in (None, ""):
        return default
    txt = str(v).strip().replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        number = float(txt)
    except ValueError:
        raise ValueError(f"Некорректное числовое значение: {v}")
    if not math.isfinite(number):
        raise ValueError("Число должно быть конечным")
    return number


def load_cost_master_xlsx(path_or_bytes: Any) -> CostMaster:
    """Load v4 cost master. Uses openpyxl at app runtime, not during artifact build."""
    from openpyxl import load_workbook  # runtime dependency
    import io

    if isinstance(path_or_bytes, (bytes, bytearray)):
        wb = load_workbook(io.BytesIO(path_or_bytes), data_only=True)
    else:
        wb = load_workbook(path_or_bytes, data_only=True)

    master = CostMaster()
    required_sheets = {"Настройки", "Ставки", "Цены", "Нормы"}
    if not required_sheets.issubset(wb.sheetnames):
        raise ValueError("Нужен мастер v4 с листами Настройки, Ставки, Цены, Нормы")

    if "Настройки" in wb.sheetnames:
        ws = wb["Настройки"]
        mapping = {
            "vat_pct": "vat_pct",
            "target_margin_pct": "target_margin_pct",
            "contingency_pct": "contingency_pct",
            "procurement_overhead_pct": "procurement_overhead_pct",
            "production_overhead_pct": "production_overhead_pct",
            "packaging_pct": "packaging_pct",
            "warranty_risk_pct": "warranty_risk_pct",
            "commissioning_reserve_pct": "commissioning_reserve_pct",
            "minimum_price_coverage_pct": "minimum_price_coverage_pct",
            "frame_steel_waste_pct": "frame_steel_waste_pct",
            "belt_length_reserve_pct": "belt_length_reserve_pct",
        }
        for row in ws.iter_rows(min_row=2, max_col=2, values_only=True):
            key = str(row[0] or "").strip()
            if key in mapping:
                setattr(master.settings, mapping[key], _f(row[1], getattr(master.settings, mapping[key])))

    if "Ставки" in wb.sheetnames:
        ws = wb["Ставки"]
        for row in ws.iter_rows(min_row=2, max_col=4, values_only=True):
            operation = str(row[0] or "").strip()
            rate = _f(row[1])
            if operation:
                master.labor_rates[operation] = LaborRate(operation, rate, str(row[2] or ""), str(row[3] or ""))

    if "Цены" in wb.sheetnames:
        ws = wb["Цены"]
        for row in ws.iter_rows(min_row=2, max_col=9, values_only=True):
            category = str(row[0] or "").strip()
            desc = str(row[1] or "").strip()
            if not category or not desc:
                continue
            master.prices.append(PriceEvidence(
                category=category,
                description=desc,
                unit=str(row[2] or "").strip(),
                unit_price_rub=_f(row[3]),
                source_type=str(row[4] or "missing").strip() or "missing",
                source_name=str(row[5] or "").strip(),
                source_url=str(row[6] or "").strip(),
                price_date=str(row[7] or "").strip(),
                note=str(row[8] or "").strip(),
            ))

    if "Нормы" in wb.sheetnames:
        ws = wb["Нормы"]
        for row in ws.iter_rows(min_row=2, max_col=2, values_only=True):
            key = str(row[0] or "").strip()
            if key:
                master.norms[key] = _f(row[1])

    validate_master(master)
    wb.close()
    return master


def validate_master(master: CostMaster) -> None:
    for key, value in asdict(master.settings).items():
        n = _f(value)
        if n < 0 or (key.endswith("_pct") and n > 100):
            raise ValueError(f"Недопустимая настройка {key}: {value}")
    if master.settings.target_margin_pct >= 100:
        raise ValueError("Рентабельность должна быть меньше 100%")
    for p in master.prices:
        if _f(p.unit_price_rub) < 0:
            raise ValueError(f"Отрицательная цена: {p.description}")
    for rate in master.labor_rates.values():
        if _f(rate.rate_rub_h) < 0:
            raise ValueError(f"Отрицательная ставка: {rate.operation}")
    for key, value in master.norms.items():
        if _f(value) < 0:
            raise ValueError(f"Отрицательная норма: {key}")


def item_key(item: Dict[str, Any]) -> str:
    return " | ".join(str(item.get(k, "")).strip() for k in ("Позиция", "Параметр", "Ед."))


def normalized_unit(unit: str) -> str:
    text = str(unit).lower().strip().rstrip(".")
    return {"штук": "шт", "шт": "шт", "комплект": "компл", "компл": "компл",
            "метр": "м", "м.п": "м", "пог.м": "м", "килограмм": "кг"}.get(text, text)


# ---------------------------------------------------------------------------
# BOM interpretation and price matching
# ---------------------------------------------------------------------------


def infer_category(position: str) -> str:
    p = position.lower()
    if "лента" in p:
        return "belt"
    if "роликоопор" in p:
        return "idler_set"
    if "ролик" in p:
        return "roller"
    if "барабан" in p:
        return "drum"
    if "мотор-редукт" in p:
        return "gearmotor"
    if "электродвиг" in p:
        return "motor"
    if "редукт" in p or "передача" in p:
        return "gearbox"
    if "балка" in p or "рама" in p:
        return "frame_steel"
    return "other"


def _numbers(text: str) -> List[float]:
    return [float(x.replace(",", ".")) for x in re.findall(r"\d+(?:[.,]\d+)?", text or "")]


def _parse_rect_tube(section: str) -> Optional[Tuple[float, float, float]]:
    # Accepts 60×40×3 / 60x40x3 / 60*40*3
    m = re.search(r"(\d+(?:[.,]\d+)?)\s*[×xхX*]\s*(\d+(?:[.,]\d+)?)\s*[×xхX*]\s*(\d+(?:[.,]\d+)?)", section or "")
    if not m:
        return None
    return tuple(float(x.replace(",", ".")) for x in m.groups())  # type: ignore


def rect_tube_mass_kg_m(section: str, density: float = 7850.0) -> Optional[float]:
    parsed = _parse_rect_tube(section)
    if not parsed:
        return None
    h, b, t = parsed
    if min(h, b, t) <= 0 or 2*t >= min(h, b):
        return None
    area_mm2 = h*b - (h-2*t)*(b-2*t)
    return area_mm2 * 1e-6 * density


def _score_candidate(item: Dict[str, Any], ev: PriceEvidence) -> float:
    """Heuristic spec match. Exact corporate item descriptions rank highest."""
    position = str(item.get("Позиция", ""))
    param = str(item.get("Параметр", ""))
    target = f"{position} {param}".lower()
    desc = ev.description.lower()
    score = 0.0
    if desc and desc in target:
        score += 100
    tnums = _numbers(target)
    dnums = _numbers(desc)
    for n in dnums:
        if any(abs(n - x) <= max(1.0, abs(n)*0.01) for x in tnums):
            score += 10
    if ev.source_type == "corporate":
        score += 5
    elif ev.source_type in {"supplier_quote", "last_purchase"}:
        score += 4
    elif ev.source_type == "online_market":
        score += 2
    return score


def resolve_price(item: Dict[str, Any], master: CostMaster) -> Optional[PriceEvidence]:
    category = infer_category(str(item.get("Позиция", "")))
    candidates = master.price_candidates(category)
    if not candidates:
        return None
    unit = normalized_unit(item.get("Ед.", ""))
    target = f"{item.get('Позиция', '')} {item.get('Параметр', '')}".lower()
    tnums = _numbers(target)
    candidates = [p for p in candidates if normalized_unit(p.unit) == unit
                  and p.source_type in SOURCE_PRIORITY
                  and (not p.item_key or p.item_key == item_key(item))
                  and all(any(abs(n - t) <= max(0.01, abs(n)*0.01) for t in tnums)
                          for n in _numbers(p.description))
                  and not ("приводн" in p.description.lower() and "хвостов" in target)
                  and not ("хвостов" in p.description.lower() and "приводн" in target)]
    if not candidates:
        return None
    best = min(candidates, key=lambda p: (SOURCE_PRIORITY[p.source_type], -_score_candidate(item, p)))
    return best if best.unit_price_rub > 0 else None


def _fallback_frame_price(item: Dict[str, Any], master: CostMaster) -> Optional[PriceEvidence]:
    steel_candidates = [p for p in master.price_candidates("steel_kg")
                        if normalized_unit(p.unit) == "кг" and p.source_type in SOURCE_PRIORITY]
    if not steel_candidates:
        return None
    mass = rect_tube_mass_kg_m(str(item.get("Параметр", "")))
    if mass is None:
        return None
    steel = min(steel_candidates, key=lambda p: SOURCE_PRIORITY[p.source_type])
    price_m = mass * steel.unit_price_rub * (1.0 + master.settings.frame_steel_waste_pct / 100.0)
    return PriceEvidence(
        category="frame_steel",
        description=f"Расчет из массы профиля {item.get('Параметр','')} × цена стали",
        unit="м",
        unit_price_rub=price_m,
        source_type="formula_estimate",
        source_name=steel.source_name or "цена стали из справочника",
        source_url=steel.source_url,
        price_date=steel.price_date,
        note=f"База {steel.unit_price_rub:.2f} руб/кг; отход {master.settings.frame_steel_waste_pct:.1f}%",
    )


def price_bom(bom: List[Dict[str, Any]], master: CostMaster) -> Tuple[List[PricedBOMItem], List[str]]:
    priced: List[PricedBOMItem] = []
    warnings: List[str] = []
    for item in bom:
        position = str(item.get("Позиция", ""))
        qty = _f(item.get("Кол-во"))
        if qty <= 0:
            raise ValueError(f"Количество должно быть больше нуля: {position}")
        unit = str(item.get("Ед.", ""))
        param = str(item.get("Параметр", ""))
        cat = infer_category(position)
        manual = item.get("Ручная цена")
        ev = PriceEvidence(**manual) if manual else resolve_price(item, master)
        if ev is not None and (not math.isfinite(ev.unit_price_rub) or ev.unit_price_rub < 0):
            raise ValueError(f"Некорректная цена: {position}")
        if ev is not None and ev.unit_price_rub == 0:
            ev = None
        if ev is None and cat == "frame_steel" and not manual:
            ev = _fallback_frame_price(item, master)
        if ev is None:
            priced.append(PricedBOMItem(
                position, qty, unit, param, cat, 0.0, 0.0,
                "missing", "", "", 0.0, 50.0, "нет цены",
                "Нужно получить цену из корпоративного справочника, поставщика или онлайн-рынка.",
            ))
            warnings.append(f"Нет цены: {position} — {param}")
            continue
        total = qty * ev.unit_price_rub
        priced.append(PricedBOMItem(
            position=position,
            quantity=qty,
            unit=unit,
            parameter=param,
            category=cat,
            unit_price_rub=ev.unit_price_rub,
            total_rub=total,
            source_type=ev.source_type,
            source_name=ev.source_name,
            source_url=ev.source_url,
            confidence=ev.confidence,
            uncertainty_pct=SOURCE_UNCERTAINTY.get(ev.source_type, 0.25) * 100.0,
            status="оценено" if ev.source_type in {"online_market", "formula_estimate"} else "цена подтверждена",
            note=ev.note,
            price_date=ev.price_date,
        ))
    return priced, warnings


# ---------------------------------------------------------------------------
# Labor model
# ---------------------------------------------------------------------------


def _norm(master: CostMaster, key: str, fallback: float) -> float:
    return master.norms.get(key, fallback)


def _rate(master: CostMaster, operation: str, general_override: float = 0.0) -> LaborRate:
    lr = master.labor_rates.get(operation)
    if lr and (lr.rate_rub_h > 0 or lr.source == "Ручная ставка операции"):
        return lr
    if general_override > 0:
        return LaborRate(operation, general_override, "ручная общая ставка", "временная ставка для предварительной оценки")
    general = master.labor_rates.get("Общая производственная")
    if general and general.rate_rub_h > 0:
        return LaborRate(operation, general.rate_rub_h, general.source or "общая ставка", general.note)
    return LaborRate(operation, 0.0, "", "ставка не заполнена")


def labor_model(product_type: str, engineering: Dict[str, Any], bom: List[Dict[str, Any]], master: CostMaster, general_rate_override: float = 0.0) -> List[LaborLine]:
    lines: List[Tuple[str, float]] = []
    if product_type == "Ленточный конвейер":
        inp = engineering.get("input", {})
        length = _f(inp.get("length_m"))
        carry_spacing = max(_f(inp.get("carrying_idler_spacing_m"), 1.2), 0.1)
        roller_count = 0.0
        belt_length = sum(_f(item.get("Кол-во")) for item in bom if infer_category(str(item.get("Позиция", ""))) == "belt")
        for item in bom:
            if infer_category(str(item.get("Позиция", ""))) in {"roller", "idler_set"}:
                roller_count += _f(item.get("Кол-во")) * (3.0 if infer_category(str(item.get("Позиция", ""))) == "idler_set" else 1.0)
        lines = [
            ("Заготовительные работы", length * _norm(master, "belt_cut_h_m", 0.22)),
            ("Сварка рамы", length * _norm(master, "belt_weld_h_m", 0.45)),
            ("Слесарная сборка", length * _norm(master, "belt_assembly_h_m", 0.35)),
            ("Установка роликов", roller_count * _norm(master, "belt_roller_h_pc", 0.10)),
            ("Приводная станция", _norm(master, "belt_drive_h", 6.0)),
            ("Натяжная станция", _norm(master, "belt_tension_h", 4.0)),
            ("Монтаж и центровка ленты", belt_length * _norm(master, "belt_install_h_m", 0.12)),
            ("Покраска", length * _norm(master, "belt_paint_h_m", 0.12)),
            ("Электромонтаж", _norm(master, "belt_electrical_base_h", 2.0) + length * _norm(master, "belt_electrical_h_m", 0.05)),
            ("Пусконаладка", _norm(master, "belt_commission_base_h", 2.0) + length * _norm(master, "belt_commission_h_m", 0.08)),
        ]
    else:
        inp = engineering.get("input", {})
        rr = engineering.get("roller", {})
        length = _f(inp.get("conveyor_length_m"))
        roller_count = _f(rr.get("roller_count"))
        driven = str(inp.get("conveyor_type", "")).lower().startswith("прив")
        lines = [
            ("Заготовительные работы", length * _norm(master, "roller_cut_h_m", 0.20)),
            ("Сварка рамы", length * _norm(master, "roller_weld_h_m", 0.35)),
            ("Слесарная сборка", length * _norm(master, "roller_assembly_h_m", 0.30)),
            ("Установка роликов", roller_count * _norm(master, "roller_install_h_pc", 0.12)),
            ("Монтаж привода", _norm(master, "roller_drive_h", 4.0) if driven else 0.0),
            ("Покраска", length * _norm(master, "roller_paint_h_m", 0.10)),
            ("Электромонтаж", _norm(master, "roller_electrical_h", 1.5) if driven else 0.0),
            ("Регулировка/испытание", length * _norm(master, "roller_adjustment_h_m", 0.15)),
        ]

    out: List[LaborLine] = []
    for op, hours in lines:
        if hours <= 0:
            continue
        lr = _rate(master, op, general_rate_override)
        out.append(LaborLine(
            operation=op,
            hours=hours,
            rate_rub_h=lr.rate_rub_h,
            cost_rub=hours * lr.rate_rub_h,
            source=lr.source,
            status="рассчитано" if lr.rate_rub_h > 0 else "нет ставки",
        ))
    return out


# ---------------------------------------------------------------------------
# Full preliminary costing
# ---------------------------------------------------------------------------


def calculate_cost(
    product_type: str,
    engineering: Dict[str, Any],
    bom: List[Dict[str, Any]],
    master: CostMaster,
    general_rate_override: float = 0.0,
    complexity_factor: float = 1.0,
) -> CostSummary:
    validate_master(master)
    if _f(general_rate_override) < 0 or _f(complexity_factor) <= 0:
        raise ValueError("Ставка должна быть неотрицательной, коэффициент сложности — положительным")
    priced, warnings = price_bom(bom, master)
    labor = labor_model(product_type, engineering, bom, master, general_rate_override)
    for line in labor:
        line.hours *= complexity_factor
        line.cost_rub = line.hours * line.rate_rub_h

    materials = sum(x.total_rub for x in priced)
    labor_cost = sum(x.cost_rub for x in labor)

    priceable_qty = sum(max(x.total_rub, 0.0) for x in priced)
    total_bom_rows = len(priced)
    priced_rows = sum(1 for x in priced if x.unit_price_rub > 0)
    coverage = (priced_rows / total_bom_rows * 100.0) if total_bom_rows else 0.0

    priced_cost_sum = sum(x.total_rub for x in priced if x.total_rub > 0)
    if priced_cost_sum > 0:
        weighted_conf = sum(x.total_rub * x.confidence for x in priced if x.total_rub > 0) / priced_cost_sum
        weighted_u = sum(x.total_rub * (x.uncertainty_pct / 100.0) for x in priced if x.total_rub > 0) / priced_cost_sum
    else:
        weighted_conf, weighted_u = 0.0, 0.30

    missing_labor = [x.operation for x in labor if x.rate_rub_h <= 0]
    if not labor:
        missing_labor.append("нормы трудоёмкости: все операции имеют нулевые часы")
    if missing_labor:
        warnings.append("Не заполнены ставки: " + ", ".join(missing_labor))

    settings = master.settings
    procurement_oh = materials * settings.procurement_overhead_pct / 100.0
    base_prod = materials + labor_cost + procurement_oh
    production_oh = labor_cost * settings.production_overhead_pct / 100.0
    subtotal = base_prod + production_oh
    contingency = subtotal * settings.contingency_pct / 100.0
    packaging = subtotal * settings.packaging_pct / 100.0
    warranty = subtotal * settings.warranty_risk_pct / 100.0
    commissioning = subtotal * settings.commissioning_reserve_pct / 100.0
    full_cost = subtotal + contingency + packaging + warranty + commissioning

    margin = settings.target_margin_pct / 100.0
    selling_ex_vat = full_cost / (1.0 - margin) if margin < 1 else full_cost
    vat = selling_ex_vat * settings.vat_pct / 100.0
    selling_inc_vat = selling_ex_vat + vat

    # Engineering-stage uncertainty is combined with evidence uncertainty.
    stage_u = 0.05
    missing_penalty = max(0.0, (100.0 - coverage) / 100.0) * 0.25
    labor_penalty = 0.08 if missing_labor else 0.0
    total_u = min(0.60, weighted_u + stage_u + missing_penalty + labor_penalty)
    low = selling_inc_vat * max(0.5, 1.0 - total_u)
    high = selling_inc_vat * (1.0 + total_u)

    status = "рассчитана стоимость введённого состава"
    if coverage < settings.minimum_price_coverage_pct or missing_labor:
        status = "предварительная оценка — требуется дозаполнение себестоимостей"
    if coverage < 40:
        status = "недостаточно цен для надежной оценки"
    from questionnaire_engine_v3 import engineering_blockers
    blockers = engineering_blockers(engineering)
    if blockers:
        status = "требуется инженерная проверка"
        warnings.extend(blockers)
    if coverage < 100 or missing_labor:
        warnings.append("Итог учитывает только заполненные цены и ставки. Это неполная сумма, а не стоимость всего оборудования; диапазон не оценивает отсутствующие позиции.")
    warnings.append("Нормы труда — предварительные; ставки и закупочные цены требуют подтверждения. Все цены в калькуляции принимаются без входного НДС.")

    if total_bom_rows and coverage < 100:
        warnings.append(f"Покрытие BOM ценами: {coverage:.0f}% ({priced_rows}/{total_bom_rows} позиций).")
    warnings.append("Результат является ориентировочным ТКП, а не окончательной калькуляцией производства/договора.")

    return CostSummary(
        product_type=product_type,
        direct_materials_rub=materials,
        direct_labor_rub=labor_cost,
        procurement_overhead_rub=procurement_oh,
        production_overhead_rub=production_oh,
        contingency_rub=contingency,
        packaging_rub=packaging,
        warranty_risk_rub=warranty,
        commissioning_reserve_rub=commissioning,
        full_cost_rub=full_cost,
        selling_price_ex_vat_rub=selling_ex_vat,
        vat_rub=vat,
        selling_price_inc_vat_rub=selling_inc_vat,
        price_low_inc_vat_rub=low,
        price_high_inc_vat_rub=high,
        price_coverage_pct=coverage,
        weighted_confidence=weighted_conf,
        weighted_uncertainty_pct=total_u * 100.0,
        pricing_status=status,
        warnings=warnings,
        priced_bom=[asdict(x) for x in priced],
        labor=[asdict(x) for x in labor],
    )


def merge_online_prices(master: CostMaster, online_rows: Iterable[Dict[str, Any]]) -> CostMaster:
    """Add market price evidence returned by the optional online price agent."""
    for row in online_rows:
        try:
            price = _f(row.get("unit_price_rub"))
            price_date = str(row.get("price_date", ""))
            from urllib.parse import urlparse
            parsed_url = urlparse(str(row.get("source_url", "")))
            if parsed_url.scheme not in {"http", "https"} or not parsed_url.netloc:
                continue
            if date.fromisoformat(price_date) > date.today():
                continue
            included = row.get("vat_included")
            vat_rate = _f(row.get("vat_pct"))
            if not isinstance(included, bool) or not 0 <= vat_rate <= 100:
                continue
            if included:
                price /= 1 + vat_rate / 100
        except (ValueError, TypeError):
            continue
        if price <= 0:
            continue
        master.prices.append(PriceEvidence(
            category=str(row.get("category", "other")),
            description=str(row.get("description", "")),
            unit=str(row.get("unit", "шт")),
            unit_price_rub=price,
            source_type="online_market",
            source_name=str(row.get("source_name", "онлайн-рынок")),
            source_url=str(row.get("source_url", "")),
            price_date=price_date,
            note=str(row.get("note", "")) + " | Нормализовано без входного НДС; проверить закупщиком",
            item_key=str(row.get("item_key", "")),
        ))
    return master
