from pathlib import Path
import sys, json
from openpyxl import load_workbook
path=next(Path(r'D:\CodexProjects\шнековый транспортер').glob('*.xlsx'))
w=load_workbook(path,data_only=False)
v=load_workbook(path,data_only=True)
for name in sys.argv[1:]:
    print('\nSHEET',name)
    for row in w[name]:
        cells=[f'{c.column_letter}: {c.value}' + (f' => {v[name][c.coordinate].value}' if c.data_type=='f' else '') for c in row if c.value is not None]
        if cells: print(str(row[0].row)+' | '+' | '.join(cells))
