"""Regression checks for the inherited v1-v4 failure modes and local web UI."""
import copy
from dataclasses import asdict
from datetime import date
import io
import json
import math
from pathlib import Path
import sys
import unittest

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE / ".deps"))
from questionnaire_parser_v3 import ExtractedField, ExtractionResult, extract_local, _normalize_value, _parse_number, _parse_bool
from questionnaire_schema_v3 import FIELD_MAP
from questionnaire_engine_v3 import calculate_from_extraction, engineering_blockers
from transporter_core_v2 import RollerInput, calc_roller
from cost_engine_v4 import (CostMaster, PriceEvidence, LaborRate, resolve_price, calculate_cost, load_cost_master_xlsx,
                            merge_online_prices, price_bom, item_key)
from web_support import demo_extraction, export_xlsx, project_bytes, load_project


def belt_package():
    return calculate_from_extraction(demo_extraction("Ленточный конвейер"))


def filled_master(package):
    master = CostMaster()
    from cost_engine_v4 import infer_category
    for item in package.bom:
        master.prices.append(PriceEvidence(infer_category(item["Позиция"]), item["Позиция"] + " " + item["Параметр"],
            item["Ед."], 1000.0, "corporate", "Тестовый источник", price_date=date.today().isoformat(), item_key=item_key(item)))
    master.labor_rates["Общая производственная"] = LaborRate("Общая производственная", 1000, "test")
    return master


class EngineeringTests(unittest.TestCase):
    def test_workbook_sample(self):
        p = calculate_from_extraction(extract_local((BASE / 'questionnaire_sample_belt_v3.xlsx').read_bytes(), 'sample.xlsx'))
        self.assertEqual(p.status, "calculated")
        self.assertFalse(engineering_blockers(p.engineering))

    def test_missing_fields_not_defaulted(self):
        ex = demo_extraction("Ленточный конвейер")
        del ex.fields['capacity_tph']
        self.assertEqual(calculate_from_extraction(ex).status, 'needs_input')

    def test_invalid_numbers_reported(self):
        for value in [-1, 0, float('nan'), float('inf'), 'text']:
            with self.subTest(value=value):
                ex = demo_extraction("Ленточный конвейер")
                ex.fields['length_m'].value = value
                p = calculate_from_extraction(ex)
                self.assertEqual(p.status, 'needs_input')
                self.assertFalse(p.bom)

    def test_zero_incline_is_valid(self):
        ex = demo_extraction("Ленточный конвейер")
        ex.fields['incline_deg'].value = 0
        self.assertEqual(calculate_from_extraction(ex).status, 'calculated')

    def test_unsupported_incline(self):
        ex = demo_extraction("Ленточный конвейер")
        ex.fields['incline_deg'].value = 80
        self.assertEqual(calculate_from_extraction(ex).status, 'needs_input')

    def test_roller_requires_length(self):
        ex = demo_extraction("Рольганг")
        del ex.fields['conveyor_length_m']
        self.assertEqual(calculate_from_extraction(ex).status, 'needs_input')

    def test_cylinder_requires_axial_length(self):
        ex = demo_extraction("Рольганг")
        ex.fields['cargo_shape'].value = 'Цилиндрический груз'
        ex.fields['cylinder_diameter_mm'] = ExtractedField('cylinder_diameter_mm', 600.0)
        self.assertEqual(calculate_from_extraction(ex).status, 'needs_input')
        ex.fields['cylinder_axial_length_mm'] = ExtractedField('cylinder_axial_length_mm', 500.0)
        self.assertEqual(calculate_from_extraction(ex).status, 'calculated')

    def test_three_rollers_and_length(self):
        ex = demo_extraction("Рольганг")
        p = calculate_from_extraction(ex)
        r = p.engineering['roller']
        self.assertLessEqual(r['selected_pitch_mm'], 800 / 3)
        self.assertGreater(r['roller_count'], 3)
        ex.fields['conveyor_length_m'].value = 12
        p2 = calculate_from_extraction(ex)
        self.assertGreater(p2.engineering['roller']['roller_count'], r['roller_count'])
        self.assertGreater(p2.engineering['roller']['simultaneous_loads'], r['simultaneous_loads'])

    def test_heavy_cargo_no_unbound_local(self):
        ex = demo_extraction("Рольганг")
        ex.fields['cargo_mass_kg'].value = 1000000
        p = calculate_from_extraction(ex)
        self.assertEqual(p.status, 'calculated')
        self.assertIsNone(p.engineering['roller']['selected_diameter_mm'])
        self.assertGreater(p.engineering['roller']['simultaneous_loads'], 1)
        self.assertTrue(engineering_blockers(p.engineering))

    def test_driven_zero_speed_rejected(self):
        ex = demo_extraction("Рольганг")
        ex.fields['speed_mps'].value = 0
        self.assertEqual(calculate_from_extraction(ex).status, 'needs_input')
        ex.fields['conveyor_type'].value = 'Гравитационный/неприводной'
        p = calculate_from_extraction(ex)
        self.assertEqual(p.status, 'calculated')
        self.assertEqual(len(p.bom), 2)


class ParserTests(unittest.TestCase):
    def test_ai_file_and_image_request_types(self):
        from unittest.mock import patch, MagicMock
        from questionnaire_parser_v3 import extract_with_openai
        for filename, purpose, content_type in [('photo.png', 'vision', 'input_image'), ('doc.pdf', 'user_data', 'input_file')]:
            module = MagicMock()
            client = module.OpenAI.return_value
            client.files.create.return_value.id = 'file-test'
            client.responses.create.return_value.output_text = '{"product_type":null,"fields":[],"warnings":[]}'
            with patch.dict(sys.modules, {'openai': module}):
                extract_with_openai(b'test', filename, 'test-key', 'test-model')
            self.assertEqual(client.files.create.call_args.kwargs['purpose'], purpose)
            self.assertEqual(client.responses.create.call_args.kwargs['input'][0]['content'][1]['type'], content_type)
            client.files.delete.assert_called_once_with('file-test')

    def test_units(self):
        cases = [('length_m', '7000', 'мм', 7), ('bulk_density_t_m3', '1 400', 'кг/м³', 1.4),
                 ('speed_mps', '18', 'м/мин', .3), ('capacity_tph', '25000', 'кг/ч', 25),
                 ('cargo_width_mm', '.5', 'м', 500)]
        for key, val, unit, result in cases:
            self.assertAlmostEqual(_normalize_value(FIELD_MAP[key], val, unit), result)

    def test_ranges_not_silently_first_number(self):
        for value in ['10-20', 'nan', 'inf', 'abc12', '12x20']:
            self.assertIsNone(_parse_number(value))
        self.assertEqual(_parse_number('1 250,5'), 1250.5)
        self.assertEqual(_parse_number('1e3'), 1000)

    def test_negative_fragility(self):
        self.assertFalse(_parse_bool('не хрупкий'))

    def test_multiline_comma_csv(self):
        data = b'product_type,conveyor\nlength_m,7\ncapacity_tph,100\n'
        ex = extract_local(data, 'data.csv')
        self.assertEqual(ex.fields['length_m'].value, 7)


class CostTests(unittest.TestCase):
    def test_zero_labor_norms_are_incomplete(self):
        p = belt_package()
        master = filled_master(p)
        master.norms = {k: 0 for k in load_cost_master_xlsx(BASE / 'cost_master_transporters_v4.xlsx').norms}
        c = calculate_cost(p.product_type, p.engineering, p.bom, master)
        self.assertFalse(c.labor)
        self.assertIn('дозаполнение', c.pricing_status)
        from web_support import cost_structure
        self.assertFalse(any('Цена' in row['Статья'] or 'цена' in row['Статья'] for row in cost_structure(c)))

    def test_portable_master(self):
        master = load_cost_master_xlsx(BASE / 'cost_master_transporters_v4.xlsx')
        self.assertIn('belt_weld_h_m', master.norms)
        self.assertTrue(all(r.rate_rub_h == 0 for r in master.labor_rates.values()))

    def test_wrong_workbook_rejected(self):
        with self.assertRaises(ValueError):
            load_cost_master_xlsx(BASE / 'questionnaire_sample_belt_v3.xlsx')

    def test_source_priority(self):
        item = {'Позиция': 'Лента конвейерная', 'Параметр': 'B=650 мм', 'Ед.': 'м'}
        master = CostMaster(prices=[PriceEvidence('belt', 'Лента конвейерная B=650 мм', 'м', 2, 'online_market'),
                                   PriceEvidence('belt', 'Лента', 'м', 10, 'corporate')])
        self.assertEqual(resolve_price(item, master).unit_price_rub, 10)

    def test_wrong_unit_and_width_rejected(self):
        item = {'Позиция': 'Лента конвейерная', 'Параметр': 'B=650 мм', 'Ед.': 'м'}
        for description, unit in [('Лента', 'кг'), ('Лента B=500 мм', 'м')]:
            master = CostMaster(prices=[PriceEvidence('belt', description, unit, 10, 'corporate')])
            self.assertIsNone(resolve_price(item, master))

    def test_date_and_manual_zero(self):
        p = belt_package()
        m = filled_master(p)
        priced, _ = price_bom(p.bom, m)
        self.assertEqual(priced[0].price_date, date.today().isoformat())
        manual = copy.deepcopy(m.prices[0])
        manual.unit_price_rub = 0
        p.bom[0]['Ручная цена'] = asdict(manual)
        self.assertEqual(price_bom(p.bom, m)[0][0].unit_price_rub, 0)

    def test_complexity_line_total_reconciles(self):
        p = belt_package()
        c = calculate_cost(p.product_type, p.engineering, p.bom, filled_master(p), complexity_factor=1.5)
        self.assertAlmostEqual(sum(x['cost_rub'] for x in c.labor), c.direct_labor_rub)
        self.assertTrue(all(abs(x['cost_rub'] - x['hours'] * x['rate_rub_h']) < 1e-6 for x in c.labor))

    def test_margin_and_vat(self):
        p = belt_package()
        c = calculate_cost(p.product_type, p.engineering, p.bom, filled_master(p))
        self.assertEqual(c.price_coverage_pct, 100)
        self.assertAlmostEqual(c.selling_price_ex_vat_rub, c.full_cost_rub / .8)
        self.assertAlmostEqual(c.selling_price_inc_vat_rub, c.selling_price_ex_vat_rub * 1.22)

    def test_missing_costs_are_not_ready(self):
        p = belt_package()
        c = calculate_cost(p.product_type, p.engineering, p.bom, CostMaster())
        self.assertEqual(c.price_coverage_pct, 0)
        self.assertIn('недостаточно', c.pricing_status)

    def test_invalid_finance_values(self):
        for margin in [100, -1, float('nan')]:
            p = belt_package()
            m = filled_master(p)
            m.settings.target_margin_pct = margin
            with self.assertRaises(ValueError):
                calculate_cost(p.product_type, p.engineering, p.bom, m)

    def test_online_tax_source_date(self):
        row = dict(category='motor', description='Электродвигатель', unit='шт', unit_price_rub=1220,
                   source_url='https://example.com/product', source_name='test', price_date=date.today().isoformat(),
                   vat_included=True, vat_pct=22)
        m = merge_online_prices(CostMaster(), [row])
        self.assertAlmostEqual(m.prices[0].unit_price_rub, 1000)
        for changes in [dict(source_url=''), dict(price_date=''), dict(vat_included=None), dict(price_date='2999-01-01')]:
            self.assertFalse(merge_online_prices(CostMaster(), [row | changes]).prices)

    def test_excel_export_and_formula_text(self):
        from openpyxl import load_workbook
        p = belt_package()
        p.audit[0]['Значение'] = '=1+1'
        c = calculate_cost(p.product_type, p.engineering, p.bom, filled_master(p))
        wb = load_workbook(io.BytesIO(export_xlsx(p, c)))
        self.assertEqual(len(wb.sheetnames), 5)
        self.assertTrue(any(x.value == '=1+1' and x.data_type == 's' for row in wb['Исходные данные'] for x in row))

    def test_saved_project_roundtrip(self):
        ex = demo_extraction('Рольганг')
        raw = project_bytes(ex, CostMaster().settings, 1200, 1.2, [], [])
        restored, obj = load_project(raw)
        self.assertEqual(restored.values(), ex.values())
        self.assertEqual(obj['rate'], 1200)

    def test_snapshot_preserves_cost_master(self):
        p = belt_package()
        master = filled_master(p)
        master.norms['belt_weld_h_m'] = 2.3
        data = project_bytes(demo_extraction('Ленточный конвейер'), master.settings, 0, 1.5, [], [], master)
        ex, obj = load_project(data)
        from cost_engine_v4 import CostSettings
        value = obj['master']
        restored = CostMaster(CostSettings(**value['settings']), {k: LaborRate(**r) for k, r in value['labor_rates'].items()},
                              [PriceEvidence(**r) for r in value['prices']], value['norms'])
        p2 = calculate_from_extraction(ex)
        c1 = calculate_cost(p.product_type, p.engineering, p.bom, master, complexity_factor=1.5)
        c2 = calculate_cost(p2.product_type, p2.engineering, p2.bom, restored, complexity_factor=obj['complexity'])
        self.assertEqual(asdict(c1), asdict(c2))


class UITests(unittest.TestCase):
    def test_complete_cost_in_ui(self):
        from streamlit.testing.v1 import AppTest
        app = AppTest.from_file(str(BASE / 'transporter_app_local.py'), default_timeout=40)
        p = belt_package()
        app.session_state['source'] = demo_extraction('Ленточный конвейер')
        app.session_state['saved'] = {'master': asdict(filled_master(p))}
        app.session_state['kind'] = 'Ленточный конвейер'
        app.run()
        self.assertFalse(app.exception)
        self.assertNotIn('Не определена', [m.value for m in app.metric])
        self.assertTrue(any('Все строки' in x.value for x in app.success))

    def test_manual_demo_invalid_reset_and_roller(self):
        from streamlit.testing.v1 import AppTest
        app = AppTest.from_file(str(BASE / 'transporter_app_local.py'), default_timeout=40).run()
        self.assertFalse(app.exception)
        app.button[1].click().run()
        self.assertFalse(app.exception)
        self.assertIn('650 мм', [m.value for m in app.metric])
        app.text_input(key='input_Ленточный конвейер_length_m').set_value('0').run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.metric), 0)
        app.button[2].click().run()
        self.assertFalse(app.exception)
        self.assertTrue(app.metric)
        app.number_input(key='cfg_rate').set_value(1250.0).run()
        self.assertFalse(app.exception)
        app.button[0].click().run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.metric), 0)


if __name__ == '__main__':
    unittest.main(verbosity=2)
