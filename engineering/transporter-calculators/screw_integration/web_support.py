"""Local UI helpers, portable exports and project snapshots. No network calls."""
from dataclasses import asdict
from datetime import date
import hashlib
import io
import json
import math

from questionnaire_parser_v3 import ExtractedField, ExtractionResult
from questionnaire_schema_v3 import FIELD_MAP


def fingerprint(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()[:16]


def demo_extraction(kind):
    values = {"project_name": "Учебный пример · " + kind, "cargo_name": "Демонстрационные данные"}
    if kind == "Ленточный конвейер":
        values.update(product_type=kind, capacity_tph=100.0, length_m=30.0, incline_deg=5.0,
                      bulk_density_t_m3=1.4, repose_angle_deg=35.0)
    else:
        values.update(product_type=kind, conveyor_length_m=6.0, cargo_shape="Прямоугольный груз",
                      cargo_length_mm=800.0, cargo_width_mm=500.0, cargo_mass_kg=100.0,
                      speed_mps=0.3, conveyor_type="Приводной")
    return ExtractionResult({k: ExtractedField(k, v, "Учебный пример", "Не данные заказчика", 1.0, "manual")
                             for k, v in values.items()}, kind, ["Учебный пример: параметры не относятся к заказу клиента."])


def load_project(data):
    obj = json.loads(data)
    if obj.get("format") != "tech-aero-local-1" or not isinstance(obj.get("fields"), dict):
        raise ValueError("Выберите проект, сохранённый этим приложением")
    fields = {k: ExtractedField(**v) for k, v in obj["fields"].items() if k in FIELD_MAP}
    from cost_engine_v4 import CostMaster, CostSettings, validate_master, _f
    validate_master(CostMaster(settings=CostSettings(**obj.get('settings', {}))))
    if not 0 <= _f(obj.get('rate', 0)) or not .1 <= _f(obj.get('complexity', 1)) <= 10:
        raise ValueError('Некорректная ставка или коэффициент в проекте')
    if obj.get('product_type') not in {'Ленточный конвейер', 'Рольганг'}:
        raise ValueError('Неизвестный тип оборудования')
    return ExtractionResult(fields, obj.get("product_type"), obj.get("warnings", []), extractor="Сохранённый проект"), obj


def project_bytes(extraction, settings, rate, complexity, price_rows, extra_rows, master=None):
    return json.dumps({"format": "tech-aero-local-1", "fields": {k: asdict(v) for k, v in extraction.fields.items()},
                       "product_type": extraction.product_type, "warnings": extraction.warnings,
                       "settings": asdict(settings), "rate": rate, "complexity": complexity,
                       "price_rows": price_rows, "extra_rows": extra_rows,
                       "master": asdict(master) if master else None},
                      ensure_ascii=False, indent=2, allow_nan=False).encode("utf-8")


BOM_LABELS = {"position": "Позиция", "quantity": "Количество", "unit": "Ед.", "parameter": "Параметр",
              "unit_price_rub": "Цена без НДС, ₽", "total_rub": "Сумма, ₽", "source_type": "Тип источника",
              "source_name": "Источник", "source_url": "Ссылка", "price_date": "Дата цены",
              "confidence": "Доверие", "status": "Статус", "note": "Примечание",
              "category": "Категория", "uncertainty_pct": "Неопределённость, %"}
LABOR_LABELS = {"operation": "Операция", "hours": "Нормо-ч", "rate_rub_h": "₽/нормо-ч",
                "cost_rub": "Стоимость, ₽", "source": "Источник ставки", "status": "Статус"}
SOURCES = {"corporate": "Корпоративная", "supplier_quote": "КП поставщика", "last_purchase": "Последняя закупка",
           "manual": "Ручное допущение", "online_market": "Онлайн-рынок", "formula_estimate": "Расчётная оценка", "missing": "Нет цены"}


def cost_structure(cost):
    complete = cost.price_coverage_pct == 100 and bool(cost.labor) and all(x['rate_rub_h'] > 0 for x in cost.labor) and cost.pricing_status != 'требуется инженерная проверка'
    rows = [{"Статья": title, "Сумма, ₽": getattr(cost, key)} for title, key in [
        ("Материалы и комплектующие", "direct_materials_rub"), ("Прямой труд", "direct_labor_rub"),
        ("Закупочные накладные", "procurement_overhead_rub"), ("Производственные накладные", "production_overhead_rub"),
        ("Резерв предварительного расчёта", "contingency_rub"), ("Упаковка", "packaging_rub"),
        ("Гарантийный резерв", "warranty_risk_rub"), ("Резерв ПНР", "commissioning_reserve_rub"),
        ("Итого по заполненному составу", "full_cost_rub"), ("Расчётная цена состава без НДС", "selling_price_ex_vat_rub"),
        ("НДС при продаже", "vat_rub"), ("Расчётная цена состава с НДС", "selling_price_inc_vat_rub")]]
    return rows if complete else rows[:-3]


def export_xlsx(package, cost):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    audit = [{**r, "Значение": str(r.get("Значение", ""))} for r in package.audit]
    sheets = {"Итоги": [{"Статья": "Статус", "Сумма, ₽": cost.pricing_status},
                         {"Статья": "Ограничение", "Сумма, ₽": "Предварительная стоимость введённого состава. Пропуски не равны нулю."},
                         {"Статья": "Покрытие ценами, %", "Сумма, ₽": cost.price_coverage_pct}] + cost_structure(cost),
              "Комплектующие": [{BOM_LABELS.get(k, k): SOURCES.get(v, v) if k == "source_type" else v for k, v in r.items()} for r in cost.priced_bom],
              "Труд": [{LABOR_LABELS.get(k, k): v for k, v in r.items()} for r in cost.labor],
              "Исходные данные": audit,
              "Допущения и замечания": [{"Замечание": x} for x in dict.fromkeys(package.assumptions + package.warnings + cost.warnings)]}
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        if not rows:
            continue
        keys = list(rows[0])
        ws.append(keys)
        for row in rows:
            ws.append([row.get(k, "") for k in keys])
            # User text stays text, including a leading equals sign.
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                elif isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00'
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="123747")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, key in enumerate(keys, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(65, max(16, len(key) + 3,
                max((len(str(r.get(key, ""))) for r in rows), default=0) + 2))
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
