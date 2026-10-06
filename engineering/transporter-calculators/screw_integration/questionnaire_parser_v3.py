"""Questionnaire ingestion for transporter calculator v3.

Two extraction paths:
1) deterministic/local: exact v3 XLSX template + simple label/value documents;
2) AI: OpenAI Responses API + Structured Outputs for arbitrary PDF/DOCX/XLSX/images.

AI is used only for extraction/normalization. Engineering formulas are executed
locally by transporter_core_v2.py after validation.
"""
from __future__ import annotations

import csv
import difflib
import io
import json
import math
import os
import re
import tempfile
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from questionnaire_schema_v3 import FIELD_MAP, FIELDS, FieldMeta


@dataclass
class ExtractedField:
    field_id: str
    value: Any
    source: str = ""
    evidence: str = ""
    confidence: float = 1.0
    origin: str = "document"  # document | default | manual


@dataclass
class ExtractionResult:
    fields: Dict[str, ExtractedField]
    product_type: Optional[str]
    warnings: List[str]
    raw_text: str = ""
    extractor: str = "local"

    def values(self) -> Dict[str, Any]:
        return {k: v.value for k, v in self.fields.items()}

    def audit_rows(self) -> List[Dict[str, Any]]:
        rows: List[Dict[str, Any]] = []
        for key, item in self.fields.items():
            meta = FIELD_MAP.get(key)
            rows.append({
                "field_id": key,
                "Параметр": meta.title if meta else key,
                "Значение": item.value,
                "Ед.": meta.unit if meta else "",
                "Источник": item.source,
                "Основание": item.evidence,
                "Доверие": item.confidence,
                "Происхождение": item.origin,
            })
        return rows


# -----------------------------------------------------------------------------
# Normalization
# -----------------------------------------------------------------------------

def _norm_text(s: Any) -> str:
    s = "" if s is None else str(s)
    s = s.lower().replace("ё", "е")
    s = re.sub(r"[^a-zа-я0-9%./+\- ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _parse_number(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(value) else None
    s = str(value).strip().replace("\xa0", " ").replace(",", ".")
    # 1 250,5 -> 1250.5
    s = re.sub(r"(?<=\d)\s+(?=\d)", "", s)
    m = re.fullmatch(r"\s*([-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?)\s*[^\d]*", s)
    if not m:
        return None
    number = float(m.group(1))
    return number if math.isfinite(number) else None


def _parse_bool(value: Any) -> Optional[bool]:
    if isinstance(value, bool):
        return value
    if value is None:
        return None
    s = _norm_text(value)
    if s in {"нет", "no", "false", "0", "-", "не", "не хрупкий", "нехрупкий"}:
        return False
    if s in {"да", "yes", "true", "1", "+", "есть", "хрупкий"}:
        return True
    if s in {"нет", "no", "false", "0", "-", "не"}:
        return False
    return None


def _enum_normalize(meta: FieldMeta, value: Any) -> Optional[str]:
    if value is None or value == "":
        return None
    s = _norm_text(value)
    if not meta.allowed:
        return str(value).strip()
    for opt in meta.allowed:
        if _norm_text(opt) == s:
            return opt
    # Product / domain-specific synonyms
    if meta.field_id == "product_type":
        if "ленточ" in s:
            return "Ленточный конвейер"
        if "рольган" in s or "роликов" in s:
            return "Рольганг"
    if meta.field_id == "cargo_shape":
        if any(x in s for x in ("цилинд", "рулон", "боч", "труб")):
            return "Цилиндрический груз"
        if any(x in s for x in ("короб", "ящик", "паллет", "прямоуг", "лист", "плита")):
            return "Прямоугольный груз"
    if meta.field_id == "conveyor_type":
        if "неприв" in s or "гравитац" in s:
            return "Гравитационный/неприводной"
        if "привод" in s:
            return "Приводной"
    if meta.field_id in {"abrasiveness", "dustiness"}:
        if "низ" in s or "мал" in s:
            return "Низкая"
        if "выс" in s or "сильн" in s:
            return "Высокая"
        if "сред" in s:
            return "Средняя"
    if meta.field_id == "service_category":
        for word, out in [("очень тяжел", "Очень тяжелые"), ("тяжел", "Тяжелые"), ("легк", "Легкие"), ("сред", "Средние")]:
            if word in s:
                return out
    if meta.field_id == "complexity":
        if "высот" in s:
            return "Высотные работы"
        if "стес" in s:
            return "Стесненные условия"
        if "цех" in s or "стандарт" in s:
            return "Стандартный цех"
    if meta.field_id == "special_execution":
        mapping = {
            "мороз": "Морозостойкая", "тепл": "Теплостойкая", "огне": "Трудновоспламеняющаяся",
            "трудновосп": "Трудновоспламеняющаяся", "пищ": "Пищевая", "общ": "Общего назначения",
        }
        for word, out in mapping.items():
            if word in s:
                return out
    # Fuzzy fallback only when quite close
    choices = {_norm_text(x): x for x in meta.allowed}
    match = difflib.get_close_matches(s, list(choices.keys()), n=1, cutoff=0.75)
    return choices[match[0]] if match else None


def _normalize_value(meta: FieldMeta, value: Any, raw_unit: str = "", raw_label: str = "") -> Any:
    if value is None or value == "":
        return None
    if meta.dtype == "str":
        return str(value).strip()
    if meta.dtype == "bool":
        return _parse_bool(value)
    if meta.dtype == "enum":
        return _enum_normalize(meta, value)
    num = _parse_number(value)
    if num is None:
        return None

    text = _norm_text(f"{raw_unit} {raw_label} {value}".replace("³", "3"))
    # Normalize to schema units for common cases.
    if meta.unit == "т/ч" and ("кг/ч" in text or "кг ч" in text):
        num /= 1000.0
    elif meta.unit == "т/м³" and ("кг/м3" in text or "кг м3" in text or "кг/м³" in text):
        num /= 1000.0
    elif meta.unit == "м" and ("мм" in text) and meta.field_id in {"length_m", "conveyor_length_m", "carry_spacing_m", "return_spacing_m", "frame_support_span_m"}:
        num /= 1000.0
    elif meta.unit == "мм" and re.search(r"(?:^| )м(?:$| )", text) and "мм" not in text:
        num *= 1000.0
    elif meta.unit == "м/с" and ("м/мин" in text or "м мин" in text):
        num /= 60.0
    return int(round(num)) if meta.dtype == "int" else float(num)


def _best_field_id(label: str) -> Tuple[Optional[str], float]:
    n = _norm_text(label)
    if not n:
        return None, 0.0
    # field_id exact support in templates
    if n in FIELD_MAP:
        return n, 1.0
    candidates: List[Tuple[float, str]] = []
    for meta in FIELDS:
        strings = (meta.title,) + meta.aliases + (meta.field_id,)
        for alias in strings:
            a = _norm_text(alias)
            if not a:
                continue
            if a == n:
                return meta.field_id, 1.0
            if a in n or n in a:
                score = min(len(a), len(n)) / max(len(a), len(n))
                score = max(score, 0.86)
            else:
                score = difflib.SequenceMatcher(None, a, n).ratio()
            candidates.append((score, meta.field_id))
    if not candidates:
        return None, 0.0
    score, field_id = max(candidates)
    return (field_id, score) if score >= 0.72 else (None, score)


# -----------------------------------------------------------------------------
# Local extraction
# -----------------------------------------------------------------------------

def _extract_xlsx(file_bytes: bytes, filename: str) -> ExtractionResult:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Для XLSX установите openpyxl") from exc

    wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    out: Dict[str, ExtractedField] = {}
    warnings: List[str] = []
    raw_lines: List[str] = []

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        # Detect v3 template by header names.
        header_idx = None
        headers: List[str] = []
        for i, row in enumerate(rows[:20]):
            hs = [_norm_text(x) for x in row]
            if "field id" in hs or "field_id" in hs or "параметр" in hs:
                if "значение" in hs or "value" in hs:
                    header_idx = i
                    headers = hs
                    break
        if header_idx is not None:
            def col(name_variants: Iterable[str]) -> Optional[int]:
                for v in name_variants:
                    nv = _norm_text(v)
                    for j, h in enumerate(headers):
                        if h == nv:
                            return j
                return None
            c_id = col(["field_id", "field id"])
            c_label = col(["Параметр", "parameter"])
            c_val = col(["Значение", "value"])
            c_unit = col(["Ед.", "Ед", "unit"])
            for row_no, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
                value = row[c_val] if c_val is not None and c_val < len(row) else None
                if value in (None, ""):
                    continue
                field_id = None
                score = 1.0
                if c_id is not None and c_id < len(row) and row[c_id]:
                    maybe = str(row[c_id]).strip()
                    if maybe in FIELD_MAP:
                        field_id = maybe
                label = str(row[c_label]).strip() if c_label is not None and c_label < len(row) and row[c_label] else ""
                if not field_id:
                    field_id, score = _best_field_id(label)
                if not field_id:
                    continue
                unit = str(row[c_unit]).strip() if c_unit is not None and c_unit < len(row) and row[c_unit] else ""
                norm = _normalize_value(FIELD_MAP[field_id], value, unit, label)
                if norm is not None:
                    out[field_id] = ExtractedField(field_id, norm, f"{filename}:{ws.title}!строка {row_no}", f"{label}: {value}", score, "document")
                raw_lines.append(f"{ws.title} | {label} | {value} | {unit}")
            continue

        # Generic sheet: inspect pairs of populated cells in the same row.
        for row_no, row in enumerate(rows, start=1):
            vals = [x for x in row if x not in (None, "")]
            if vals:
                raw_lines.append(" | ".join(str(x) for x in vals))
            if len(vals) < 2:
                continue
            label, value = vals[0], vals[1]
            field_id, score = _best_field_id(str(label))
            if field_id and score >= 0.78:
                norm = _normalize_value(FIELD_MAP[field_id], value, "", str(label))
                if norm is not None and field_id not in out:
                    out[field_id] = ExtractedField(field_id, norm, f"{filename}:{ws.title}!строка {row_no}", f"{label}: {value}", score, "document")

    return _finish_extraction(out, warnings, "\n".join(raw_lines), "local-xlsx")


def _extract_csv(file_bytes: bytes, filename: str) -> ExtractionResult:
    text = file_bytes.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text), delimiter=";"))
    if rows and max(map(len, rows)) <= 1:
        rows = list(csv.reader(io.StringIO(text), delimiter=","))
    out: Dict[str, ExtractedField] = {}
    raw_lines: List[str] = []
    for i, row in enumerate(rows, start=1):
        vals = [x.strip() for x in row if str(x).strip()]
        if not vals:
            continue
        raw_lines.append(" | ".join(vals))
        if len(vals) < 2:
            continue
        field_id = vals[0] if vals[0] in FIELD_MAP else None
        score = 1.0
        label = vals[1] if field_id and len(vals) >= 3 else vals[0]
        value = vals[2] if field_id and len(vals) >= 3 else vals[1]
        unit = vals[3] if field_id and len(vals) >= 4 else ""
        if not field_id:
            field_id, score = _best_field_id(label)
        if field_id:
            norm = _normalize_value(FIELD_MAP[field_id], value, unit, label)
            if norm is not None:
                out[field_id] = ExtractedField(field_id, norm, f"{filename}:строка {i}", f"{label}: {value}", score, "document")
    return _finish_extraction(out, [], "\n".join(raw_lines), "local-csv")


def _extract_textual(file_bytes: bytes, filename: str, suffix: str) -> ExtractionResult:
    warnings: List[str] = []
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("Для PDF установите pypdf") from exc
        reader = PdfReader(io.BytesIO(file_bytes))
        pages = []
        for i, page in enumerate(reader.pages, 1):
            txt = page.extract_text() or ""
            pages.append(f"[PAGE {i}]\n{txt}")
        text = "\n".join(pages)
        if len(text.strip()) < 50:
            warnings.append("PDF почти не содержит текстового слоя. Для скана используйте режим ИИ/vision.")
    elif suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("Для DOCX установите python-docx") from exc
        doc = Document(io.BytesIO(file_bytes))
        chunks = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(chunks)
    else:
        text = file_bytes.decode("utf-8", errors="replace")

    out: Dict[str, ExtractedField] = {}
    for line_no, line in enumerate(text.splitlines(), 1):
        line = line.strip()
        if not line:
            continue
        # label: value, label - value, or first two table cells
        parts = re.split(r"\s*(?:[:=]|\t|\|)\s*", line, maxsplit=2)
        if len(parts) < 2:
            parts = re.split(r"\s+-\s+", line, maxsplit=1)
        if len(parts) < 2:
            continue
        label, value = parts[0], parts[1]
        field_id, score = _best_field_id(label)
        if field_id and score >= 0.79:
            norm = _normalize_value(FIELD_MAP[field_id], value, "", label)
            if norm is not None and field_id not in out:
                out[field_id] = ExtractedField(field_id, norm, f"{filename}:строка {line_no}", line[:240], score, "document")
    return _finish_extraction(out, warnings, text, f"local-{suffix.lstrip('.') or 'txt'}")


def _finish_extraction(fields: Dict[str, ExtractedField], warnings: List[str], raw_text: str, extractor: str) -> ExtractionResult:
    ptype = None
    if "product_type" in fields:
        ptype = fields["product_type"].value
    else:
        belt_hits = sum(k in fields for k in ("capacity_tph", "bulk_density_t_m3", "repose_angle_deg", "incline_deg"))
        roller_hits = sum(k in fields for k in ("cargo_mass_kg", "cargo_shape", "conveyor_type", "cargo_length_mm", "cylinder_diameter_mm"))
        if belt_hits >= 2 and belt_hits > roller_hits:
            ptype = "Ленточный конвейер"
        elif roller_hits >= 2 and roller_hits > belt_hits:
            ptype = "Рольганг"
        if ptype:
            fields["product_type"] = ExtractedField("product_type", ptype, "определено по набору полей", "автодетект", 0.90, "document")
    return ExtractionResult(fields, ptype, warnings, raw_text, extractor)


def extract_local(file_bytes: bytes, filename: str) -> ExtractionResult:
    suffix = Path(filename).suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _extract_xlsx(file_bytes, filename)
    if suffix == ".csv":
        return _extract_csv(file_bytes, filename)
    if suffix in {".pdf", ".docx", ".txt", ".md"}:
        return _extract_textual(file_bytes, filename, suffix)
    raise ValueError(f"Локальный парсер не поддерживает {suffix}. Используйте XLSX/CSV/PDF/DOCX/TXT или режим ИИ.")


# -----------------------------------------------------------------------------
# AI extraction via Responses API Structured Outputs
# -----------------------------------------------------------------------------
AI_FIELD_ENUM = [f.field_id for f in FIELDS]

AI_EXTRACTION_SCHEMA: Dict[str, Any] = {
    "type": "object",
    "properties": {
        "product_type": {"type": ["string", "null"], "enum": ["Ленточный конвейер", "Рольганг", None]},
        "fields": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "field_id": {"type": "string", "enum": AI_FIELD_ENUM},
                    "value": {"type": ["string", "number", "boolean", "null"]},
                    "unit": {"type": "string"},
                    "evidence": {"type": "string"},
                    "confidence": {"type": "number", "minimum": 0, "maximum": 1},
                },
                "required": ["field_id", "value", "unit", "evidence", "confidence"],
                "additionalProperties": False,
            },
        },
        "warnings": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["product_type", "fields", "warnings"],
    "additionalProperties": False,
}

AI_INSTRUCTIONS = """Ты извлекаешь данные из листа опросного на ленточный конвейер или рольганг.
Верни только данные, которые действительно присутствуют в документе или однозначно следуют из единиц измерения.
Не вычисляй инженерные параметры, не придумывай отсутствующие значения и не подставляй типовые нормы.
Нормализуй единицы к единицам canonical schema: т/ч, м, градусы, т/м3, мм, кг, м/с, рубли.
Для каждого поля дай короткое evidence — фрагмент/подпись из документа. confidence 0..1.
Если встречаются противоречащие значения — выбери наиболее явно подписанное, а конфликт опиши в warnings.
Если тип транспортера не указан, определи его только при высокой уверенности по содержанию документа.
"""


def extract_with_openai(file_bytes: bytes, filename: str, api_key: Optional[str] = None, model: Optional[str] = None) -> ExtractionResult:
    """Extract arbitrary questionnaire with OpenAI Responses API.

    Current API path uses file upload with purpose=user_data and Structured Outputs.
    Requires OPENAI_API_KEY and an explicit model or OPENAI_MODEL.
    """
    try:
        from openai import OpenAI
    except ImportError as exc:
        raise RuntimeError("Для режима ИИ установите пакет openai") from exc

    api_key = api_key or os.environ.get("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("Не задан OPENAI_API_KEY")
    model = model or os.environ.get("OPENAI_MODEL")
    if not model:
        raise RuntimeError("Укажите модель API для распознавания")
    client = OpenAI(api_key=api_key, timeout=90.0, max_retries=1)

    uploaded = None
    tmp_path = None
    try:
        suffix = Path(filename).suffix or ".bin"
        is_image = suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            tmp.write(file_bytes)
            tmp_path = tmp.name
        with open(tmp_path, "rb") as fh:
            uploaded = client.files.create(file=fh, purpose="vision" if is_image else "user_data")

        response = client.responses.create(
            model=model,
            instructions=AI_INSTRUCTIONS,
            input=[{
                "role": "user",
                "content": [
                    {"type": "input_text", "text": "Извлеки структурированные данные из этого листа опросного в JSON по заданной схеме."},
                    {"type": "input_image" if is_image else "input_file", "file_id": uploaded.id},
                ],
            }],
            text={
                "format": {
                    "type": "json_schema",
                    "name": "transporter_questionnaire",
                    "schema": AI_EXTRACTION_SCHEMA,
                    "strict": True,
                },
            },
            store=False,
        )
        data = json.loads(response.output_text)
        out: Dict[str, ExtractedField] = {}
        warnings = list(data.get("warnings") or [])
        for item in data.get("fields", []):
            key = item.get("field_id")
            if key not in FIELD_MAP:
                continue
            raw_val = item.get("value")
            norm = _normalize_value(FIELD_MAP[key], raw_val, item.get("unit", ""), FIELD_MAP[key].title)
            if norm is None:
                continue
            out[key] = ExtractedField(
                field_id=key,
                value=norm,
                source=f"{filename} / AI:{model}",
                evidence=str(item.get("evidence") or "")[:500],
                confidence=float(item.get("confidence", 0.0)),
                origin="document",
            )
        ptype = data.get("product_type")
        if ptype and "product_type" not in out:
            out["product_type"] = ExtractedField("product_type", ptype, f"{filename} / AI:{model}", "тип определен ИИ", 0.90, "document")
        return _finish_extraction(out, warnings, "", f"openai:{model}")
    finally:
        if uploaded is not None:
            try:
                client.files.delete(uploaded.id)
            except Exception:
                pass
        if tmp_path:
            try:
                os.unlink(tmp_path)
            except OSError:
                pass


def apply_defaults(result: ExtractionResult) -> ExtractionResult:
    """Fill only declared corporate/prototype defaults; all are audit-marked."""
    fields = dict(result.fields)
    ptype = result.product_type or fields.get("product_type", ExtractedField("product_type", None)).value
    relevant = "belt" if ptype == "Ленточный конвейер" else "roller" if ptype == "Рольганг" else ""
    # Common + relevant defaults. Do not default product type or required critical fields.
    for meta in FIELDS:
        if meta.default is None or meta.field_id in fields:
            continue
        if relevant == "belt" and meta.section == "Рольганг":
            continue
        if relevant == "roller" and meta.section == "Ленточный":
            continue
        fields[meta.field_id] = ExtractedField(
            meta.field_id, meta.default, "типовое допущение прототипа", meta.note or "значение по умолчанию v3; не корпоративная норма", 1.0, "default"
        )
    return ExtractionResult(fields, ptype, list(result.warnings), result.raw_text, result.extractor)
