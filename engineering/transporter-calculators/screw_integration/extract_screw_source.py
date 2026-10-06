"""Read-only extraction; cached source values are evidence, never executable formulas."""
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

source = next(Path(r'D:\CodexProjects\шнековый транспортер').glob('*.xlsx'))
wb = load_workbook(source, data_only=True)
formulas = load_workbook(source, data_only=False)
data = {'filename': source.name, 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
        'import_date': '2026-09-06', 'sheets': {}}
for sheet in wb:
    data['sheets'][sheet.title] = {c.coordinate: {'value': c.value,
        'formula': formulas[sheet.title][c.coordinate].value if formulas[sheet.title][c.coordinate].data_type == 'f' else None}
        for row in sheet for c in row if c.value is not None}
data['operations'] = [{'name': wb['Op24_v9'][f'B{r}'].value, 'person_h': wb['Op24_v9'][f'C{r}'].value,
    'machine_h': wb['Op24_v9'][f'D{r}'].value, 'resource': wb['Op24_v9'][f'E{r}'].value,
    'source': f'Op24_v9!B{r}:M{r}'} for r in range(4,16)]
data['buy'] = [{'name': wb['Buy24_v10'][f'B{r}'].value, 'qty': wb['Buy24_v10'][f'C{r}'].value,
    'source': f'Buy24_v10!B{r}:M{r}'} for r in range(4,27)]
data['make'] = [{'designation': wb['BOM24_v10'][f'C{r}'].value, 'name': wb['BOM24_v10'][f'D{r}'].value,
    'qty': wb['BOM24_v10'][f'F{r}'].value, 'dimensions': wb['BOM24_v10'][f'G{r}'].value,
    'material': wb['BOM24_v10'][f'I{r}'].value, 'source': f'BOM24_v10!C{r}:N{r}'} for r in range(14,44)]
target = Path(__file__).parent / 'references' / 'screw_conveyors' / 'workbook_v11.json'
target.parent.mkdir(parents=True, exist_ok=True)
target.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
print(json.dumps({'snapshot': str(target), 'sha256': data['sha256'], 'operations': len(data['operations']),
                  'buy': len(data['buy']), 'make': len(data['make'])}, ensure_ascii=False))
