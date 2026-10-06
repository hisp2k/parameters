from pathlib import Path
from docx import Document
import re,json
paths=[Path(r'C:\Users\adm\AppData\Local\Temp\codex-file-preview-5k2l7S\вопросы 03.09. часть2.docx'),Path(r'C:\Users\adm\AppData\Local\Temp\codex-file-preview-4Kmfi5\Вопросы 03.09.26.docx'),Path(r'C:\Users\adm\AppData\Local\Temp\codex-file-preview-IBjb89\Вторая партия вопросов 2.09.2026.docx')]
out=[]
for path in paths:
 doc=Document(path); texts=[p.text for p in doc.paragraphs]
 for table in doc.tables:
  texts+=[' | '.join(c.text for c in row.cells) for row in table.rows]
 hits=[]
 for i,s in enumerate(texts):
  if re.search(r'NEXT|КЕ\s*50|KE\s*50|6163|0i.?TD|SINUMERIK|Siemens|токарн',s,re.I):
   hits.append({'index':i,'text':s,'next':texts[i+1:i+3]})
 out.append({'source_file':path.name,'hits':hits})
Path(r'C:\Users\adm\Documents\Codex\2026-10-04\new-chat-2\work\equipment_source_extract.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(out,ensure_ascii=False))
